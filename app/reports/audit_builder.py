"""The agent trail: what each specialist did, and what was corrected (PRD s14, s21).

Rendered from the database, with no model call. That matters twice over: the
record cannot be paraphrased on its way to the page, and this document still
ships when composition fails — which is precisely when someone needs to see what
actually happened.

It is the working record, not the deliverable. The technical report says what the
team concluded; this says how they got there, who disagreed, what the review
caught, and what was superseded on the way.
"""

from __future__ import annotations

import hashlib
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document as new_document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.audit.logger import get_logger
from app.reports.builder import DISCLAIMER, build_filename
from app.schemas.specialist import SpecialistResult

log = get_logger(__name__)

GENERATOR_VERSION = "1.0.0"
TABLE_STYLE = "Light Grid Accent 1"

BASIS_LABEL = {
    "fact": "Fact",
    "inference": "Inference",
    "assumption": "ASSUMPTION",
    "estimate": "ESTIMATE",
}

STATUS_LABEL = {
    "complete": "Complete",
    "needs_information": "NEEDS INFORMATION",
    "blocked": "BLOCKED",
    "failed": "FAILED",
    "cancelled": "Cancelled",
    "running": "Running",
    "pending": "Pending",
}

# Events worth a line in the timeline. The rest are machinery.
TIMELINE_EVENTS = {
    "run.created": "Run created",
    "plan.created": "Plan created",
    "approval.requested": "Approval requested",
    "approval.decided": "Approval decided",
    "task.dispatched": "Specialists dispatched",
    "task.complete": "Specialist returned",
    "conflict.detected": "Conflicts detected",
    "challenge.complete": "Independent challenge complete",
    "synthesis.complete": "Synthesis complete",
    "quality_gate.evaluated": "Quality gate",
    "report.generated": "Report generated",
    "run.complete": "Run complete",
    "run.review_required": "Review required",
    "run.failed": "Run failed",
}


@dataclass
class AuditInputs:
    run_ref: str
    document_title: str
    generated_at: str
    status: str
    started_at: datetime
    tasks: list[dict[str, Any]] = field(default_factory=list)
    results: list[SpecialistResult] = field(default_factory=list)
    agent_names: dict[str, str] = field(default_factory=dict)
    conflicts: list[dict[str, Any]] = field(default_factory=list)
    approvals: list[dict[str, Any]] = field(default_factory=list)
    artifacts: list[dict[str, Any]] = field(default_factory=list)
    evidence: list[dict[str, Any]] = field(default_factory=list)
    events: list[dict[str, Any]] = field(default_factory=list)
    report_version: int = 1


@dataclass
class RenderedAudit:
    path: Path
    filename: str
    sha256: str
    word_count: int


class AuditTrailRenderer:
    def __init__(self, output_dir: Path | str):
        self.output_dir = Path(output_dir).expanduser()

    # -- primitives ---------------------------------------------------------

    def _para(self, doc: DocxDocument, text: str, *, italic: bool = False, bold: bool = False):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(text)
        run.italic = italic
        run.bold = bold
        return paragraph

    def _table(self, doc: DocxDocument, columns: list[str], rows: list[list[str]]) -> None:
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = TABLE_STYLE
        for cell, header in zip(table.rows[0].cells, columns, strict=True):
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)

    def _name(self, inputs: AuditInputs, agent_id: str) -> str:
        return f"{agent_id} — {inputs.agent_names.get(agent_id, agent_id)}"

    # -- sections -----------------------------------------------------------

    def _cover(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        title = doc.add_heading("Agent Trail and Audit Record", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(inputs.document_title)
        run.bold = True
        run.font.size = Pt(12)

        if inputs.status.lower() in ("review_required", "failed"):
            warning = doc.add_paragraph()
            warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
            flag = warning.add_run(inputs.status.replace("_", " ").upper())
            flag.bold = True
            flag.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

        returned = [r for r in inputs.results]
        incomplete = [r for r in returned if r.status != "complete"]
        self._table(
            doc,
            ["Field", "Value"],
            [
                ["Run reference", inputs.run_ref],
                ["Report version", str(inputs.report_version)],
                ["Generated", inputs.generated_at],
                ["Final status", inputs.status.replace("_", " ").title()],
                ["Specialist tasks", str(len(inputs.tasks))],
                ["Results returned", str(len(returned))],
                ["Returned incomplete", str(len(incomplete))],
                ["Conflicts recorded", str(len(inputs.conflicts))],
                ["Evidence items", str(len(inputs.evidence))],
            ],
        )
        self._para(
            doc,
            "This is the working record of how the analysis was produced. The technical "
            "report carries the conclusions; this carries the route to them, including "
            "what was challenged, what was corrected and what remains unresolved.",
            italic=True,
        )
        self._para(doc, DISCLAIMER, italic=True)
        doc.add_page_break()

    def _timeline(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        doc.add_heading("1. Execution timeline", level=1)
        if not inputs.events:
            self._para(doc, "No events recorded.", italic=True)
            return
        rows = []
        for event in inputs.events:
            label = TIMELINE_EVENTS.get(event["type"])
            if not label:
                continue
            elapsed = int((event["at"] - inputs.started_at).total_seconds())
            detail = event.get("detail") or ""
            rows.append([f"{elapsed // 60}m {elapsed % 60:02d}s", label, detail[:110]])
        self._table(doc, ["Elapsed", "Event", "Detail"], rows)

    def _workstreams(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        doc.add_heading("2. Workstreams as planned and executed", level=1)
        if not inputs.tasks:
            self._para(doc, "No tasks were created.", italic=True)
            return
        self._table(
            doc,
            ["Task", "Specialist", "Group", "Status", "Objective"],
            [
                [
                    t["task_id"],
                    self._name(inputs, t["agent_id"]),
                    str(t["group"]),
                    STATUS_LABEL.get(t["status"], t["status"]),
                    t["objective"][:130],
                ]
                for t in inputs.tasks
            ],
        )
        self._para(
            doc,
            "Tasks in the same group ran in parallel; a later group waited for the one before it.",
            italic=True,
        )

    def _each_specialist(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        doc.add_heading("3. What each specialist reported", level=1)
        if not inputs.results:
            self._para(doc, "No specialist returned a result.", italic=True)
            return

        for result in inputs.results:
            doc.add_heading(f"{self._name(inputs, result.agent_id)} · {result.task_id}", level=2)
            self._para(
                doc,
                f"Status: {STATUS_LABEL.get(result.status, result.status)} · "
                f"Confidence: {result.confidence.replace('_', ' ')}",
                italic=True,
            )
            self._para(doc, result.executive_summary)

            if result.findings:
                doc.add_heading("Findings, and what each rests on", level=3)
                self._table(
                    doc,
                    ["Basis", "Finding", "Evidence"],
                    [
                        [
                            BASIS_LABEL.get(f.basis, f.basis),
                            f.statement[:400],
                            ", ".join(e.evidence_id for e in f.evidence_refs) or "—",
                        ]
                        for f in result.findings
                    ],
                )
                # s15: this is the check the evidence policy exists to enforce.
                unsupported = result.unsupported_findings()
                if unsupported:
                    self._para(
                        doc,
                        f"{len(unsupported)} finding(s) assert fact or inference without an "
                        "evidence reference. The runtime asks for a correction once; anything "
                        "still listed here was not corrected.",
                        italic=True,
                    )

            for heading, items in (
                ("Assumptions stated", result.assumptions),
                ("Information it needed and did not have", result.information_required),
                ("Recommendations", result.recommendations),
            ):
                if items:
                    doc.add_heading(heading, level=3)
                    for item in items:
                        doc.add_paragraph(item, style="List Bullet")

            if result.risks:
                doc.add_heading("Risks raised", level=3)
                self._table(
                    doc,
                    ["Risk", "Likelihood", "Impact", "Mitigation"],
                    [
                        [r.description[:200], r.likelihood, r.impact, r.mitigation or "—"]
                        for r in result.risks
                    ],
                )

            if result.conflicts_or_questions:
                doc.add_heading("Questions and disagreements it raised", level=3)
                for raised in result.conflicts_or_questions:
                    detail = raised.description[:400]
                    if raised.conflicting_with_agent:
                        detail += f" (with {raised.conflicting_with_agent})"
                    doc.add_paragraph(detail, style="List Bullet")

            if result.escalation_required:
                self._para(
                    doc,
                    f"Escalated to human review: {result.escalation_reason or 'no reason given'}",
                    bold=True,
                )
            if result.artifacts:
                self._para(doc, "Work products: " + ", ".join(result.artifacts), italic=True)

    def _challenge(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        doc.add_heading("4. What the independent challenge found", level=1)
        if not inputs.conflicts:
            self._para(
                doc,
                "No material contradiction was recorded between specialists.",
                italic=True,
            )
        else:
            self._table(
                doc,
                ["Description", "Between", "Severity", "Material", "State", "Resolution"],
                [
                    [
                        c["description"][:220],
                        ", ".join(c.get("agents") or []) or "—",
                        c.get("severity", "—"),
                        "Yes" if c.get("material") else "No",
                        c.get("state", "—"),
                        c.get("resolution") or "Not resolved",
                    ]
                    for c in inputs.conflicts
                ],
            )
            open_ones = [c for c in inputs.conflicts if c.get("state") == "open"]
            if open_ones:
                self._para(
                    doc,
                    f"{len(open_ones)} disagreement(s) remain open. They were not averaged "
                    "away: the report states them as unresolved.",
                    italic=True,
                )

        superseded = [a for a in inputs.artifacts if a.get("status") == "superseded"]
        doc.add_heading("5. What was revised after review", level=1)
        if superseded:
            self._table(
                doc,
                ["Work product", "Produced by", "Replaced by"],
                [
                    [
                        f"{a['ref']} — {a['title'][:70]}",
                        a.get("agent") or "—",
                        a.get("superseded_by") or "—",
                    ]
                    for a in superseded
                ],
            )
            self._para(
                doc,
                "A superseded draft is kept for traceability and excluded from the "
                "technical report's conclusions.",
                italic=True,
            )
        else:
            self._para(doc, "No work product was superseded during this run.", italic=True)

    def _evidence_and_decisions(self, doc: DocxDocument, inputs: AuditInputs) -> None:
        doc.add_heading("6. Evidence recorded during the run", level=1)
        if inputs.evidence:
            self._table(
                doc,
                ["ID", "Reliability", "Source", "Recorded by"],
                [
                    [
                        e["ref"],
                        str(e.get("reliability", "—")),
                        str(e.get("title", ""))[:90],
                        e.get("agent") or "—",
                    ]
                    for e in inputs.evidence
                ],
            )
        else:
            self._para(doc, "No evidence items were recorded.", italic=True)

        doc.add_heading("7. Human decisions", level=1)
        if inputs.approvals:
            self._table(
                doc,
                ["Gate", "Class", "Status", "Decided by", "When"],
                [
                    [
                        a["gate"],
                        a.get("decision_class", "—").upper(),
                        a["status"],
                        a.get("decided_by") or "—",
                        a.get("decided_at") or "—",
                    ]
                    for a in inputs.approvals
                ],
            )
        else:
            self._para(doc, "No gate required a human decision.", italic=True)

    # -- entry point --------------------------------------------------------

    def render(self, inputs: AuditInputs, *, project_name: str) -> RenderedAudit:
        doc = new_document()
        self._cover(doc, inputs)
        self._timeline(doc, inputs)
        self._workstreams(doc, inputs)
        self._each_specialist(doc, inputs)
        self._challenge(doc, inputs)
        self._evidence_and_decisions(doc, inputs)

        self.output_dir.mkdir(parents=True, exist_ok=True)
        filename = build_filename(
            project_name, "Agent_Trail_and_Audit", inputs.run_ref, inputs.report_version
        )
        path = self.output_dir / filename
        if path.exists():
            raise FileExistsError(
                f"{filename} already exists; bump the version rather than overwriting"
            )
        doc.save(str(path))
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        words = sum(len(p.text.split()) for p in doc.paragraphs)
        log.info("audit_trail_rendered", filename=filename, results=len(inputs.results))
        return RenderedAudit(path=path, filename=filename, sha256=digest, word_count=words)


async def build_audit_inputs(
    *,
    run_id: uuid.UUID,
    run_ref: str,
    document_title: str,
    generated_at: str,
    status: str,
    report_version: int,
    agent_names: dict[str, str],
) -> AuditInputs:
    """Read the run's record. Everything here is already persisted."""
    from app.storage.db import session_scope
    from app.storage.repositories import (
        ArtifactRepository,
        AuditRepository,
        EvidenceRepository,
        GovernanceRepository,
        RunRepository,
    )

    async with session_scope() as session:
        runs = RunRepository(session)
        run = await runs.get(run_id)
        rows = await runs.list_tasks(run_id)
        records = await runs.list_results(run_id)
        governance = GovernanceRepository(session)
        conflicts = await governance.list_conflicts(run_id)
        approvals = await governance.pending_approvals(run_id)
        from sqlalchemy import select

        from app.storage.models import Approval

        decided = (
            (
                await session.execute(
                    select(Approval).where(Approval.run_id == run_id, Approval.status != "pending")
                )
            )
            .scalars()
            .all()
        )
        artifacts = await ArtifactRepository(session).list_for_run(run_id)
        evidence = await EvidenceRepository(session).list_for_run(run_id)
        events = await AuditRepository(session).for_run(run_id)

    results: list[SpecialistResult] = []
    for record in records:
        if record.structured_result:
            try:
                results.append(SpecialistResult.model_validate(record.structured_result))
            except Exception:  # noqa: BLE001 - a malformed record must not lose the rest
                log.warning("audit_result_unreadable", task=record.planner_task_id)

    return AuditInputs(
        run_ref=run_ref,
        document_title=document_title,
        generated_at=generated_at,
        status=status,
        started_at=run.created_at if run else (events[0].created_at if events else datetime.now()),
        report_version=report_version,
        agent_names=agent_names,
        tasks=[
            {
                "task_id": t.planner_task_id,
                "agent_id": t.agent_id,
                "group": t.execution_group,
                "status": t.status,
                "objective": t.objective,
            }
            for t in rows
        ],
        results=results,
        conflicts=[
            {
                "description": c.description,
                "agents": c.agent_ids,
                "severity": c.severity,
                "material": c.material,
                "state": c.state,
                "resolution": c.resolution,
            }
            for c in conflicts
        ],
        approvals=[
            {
                "gate": a.gate,
                "decision_class": a.decision_class,
                "status": a.status,
                "decided_by": a.decided_by,
                "decided_at": a.decided_at.strftime("%H:%M:%S") if a.decided_at else None,
            }
            for a in [*decided, *approvals]
        ],
        artifacts=[
            {
                "ref": a.artifact_ref,
                "title": a.title,
                "status": a.status,
                "agent": a.created_by_agent,
                "superseded_by": a.superseded_by_ref,
            }
            for a in artifacts
            if a.kind == "work_product"
        ],
        evidence=[
            {
                "ref": e.evidence_ref,
                "reliability": e.reliability_level,
                "title": e.source_title,
                "agent": e.originating_agent,
            }
            for e in evidence
        ],
        events=[
            {"type": e.event_type, "at": e.created_at, "detail": _event_detail(e.event_metadata)}
            for e in events
        ],
    )


def _event_detail(metadata: dict[str, Any] | None) -> str:
    if not metadata:
        return ""
    for key in ("agent_id", "status", "gate", "reasons", "agents", "signals", "filename"):
        if key in metadata:
            value = metadata[key]
            return ", ".join(map(str, value)) if isinstance(value, list) else str(value)
    return ""
