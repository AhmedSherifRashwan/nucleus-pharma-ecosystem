"""The report is the regulated deliverable: nothing may be silently dropped (PRD s18)."""

from __future__ import annotations

import pytest
from docx import Document

from app.reports.builder import (
    ReportBuilder,
    ReportInputs,
    build_filename,
    safe_component,
)
from app.schemas.executive import DecisionRequired, ExecutiveDecisionBrief
from app.schemas.specialist import EvidenceRef, Finding, SpecialistResult

NAMES = {"A07": "Softgel Shell Formulation Expert", "A10": "Stability Studies Expert"}


def brief(**kw) -> ExecutiveDecisionBrief:
    base = {
        "objective": "Assess replacing gelatin with HPMC.",
        "executive_conclusion": "Feasible with mitigations.",
        "recommendation": "Proceed to a pilot batch.",
        "confidence": "medium",
        "consulted_agents": ["A07", "A10"],
    }
    return ExecutiveDecisionBrief(**{**base, **kw})


def result(agent_id: str = "A07", **kw) -> SpecialistResult:
    base = {
        "agent_id": agent_id,
        "task_id": "T1",
        "status": "complete",
        "executive_summary": "Shell change is workable.",
        "findings": [
            Finding(
                statement="HPMC shells are compatible with the fill.",
                basis="inference",
                evidence_refs=[EvidenceRef(evidence_id="E-1")],
            )
        ],
        "confidence": "medium",
    }
    return SpecialistResult(**{**base, **kw})


def text_of(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def build(tmp_path, **kw):
    inputs = ReportInputs(
        run_ref=kw.pop("run_ref", "RUN-1"),
        project_name=kw.pop("project_name", "Project Alpha"),
        brief=kw.pop("brief", brief()),
        agent_names=NAMES,
        **kw,
    )
    return ReportBuilder(tmp_path).build(inputs)


def test_filename_follows_the_naming_convention():
    name = build_filename("Project Alpha", "Executive_Decision_Brief", "RUN-12", 3)
    assert name == "Project_Alpha_Executive_Decision_Brief_RUN-12_v3.docx"


def test_filename_components_cannot_escape_the_output_directory():
    assert "/" not in safe_component("../../etc/passwd")
    assert safe_component("../../etc/passwd") == ".._.._etc_passwd"


def test_report_carries_conclusion_recommendation_and_disclaimer(tmp_path):
    out = build(tmp_path)
    body = text_of(out.path)
    assert "Feasible with mitigations." in body
    assert "Proceed to a pilot batch." in body
    assert "not a regulatory approval" in body


def test_every_mandated_section_is_present(tmp_path):
    body = text_of(build(tmp_path).path)
    for heading in (
        "Objective and scope",
        "Executive conclusion",
        "Recommendation",
        "Confidence and evidence quality",
        "Specialist workstreams and findings",
        "Implications",
        "Key evidence and provenance register",
        "Risks and mitigations",
        "Agent disagreements and unresolved conflicts",
        "Critical assumptions",
        "Information still required",
        "Recommended next actions",
        "Decisions required from you",
        "Appendix",
    ):
        assert heading in body, f"missing section: {heading}"


def test_empty_sections_say_so_rather_than_disappearing(tmp_path):
    body = text_of(build(tmp_path).path)
    assert "None identified." in body


def test_disagreements_reach_the_reader(tmp_path):
    out = build(
        tmp_path,
        brief=brief(agent_disagreements=["A07 and A10 disagree on the shelf-life driver."]),
    )
    assert "A07 and A10 disagree on the shelf-life driver." in text_of(out.path)


def test_findings_are_labelled_with_their_basis(tmp_path):
    out = build(tmp_path, results=[result()])
    body = text_of(out.path)
    assert "(inference)" in body
    assert "E-1" in body


def test_review_required_is_stamped_on_the_title_page(tmp_path):
    out = build(
        tmp_path,
        brief=brief(review_required=True, review_reason="A specialist failed to return."),
    )
    body = text_of(out.path)
    assert "REVIEW REQUIRED" in body
    assert "A specialist failed to return." in body


def test_decisions_required_are_listed_with_their_class(tmp_path):
    out = build(
        tmp_path,
        brief=brief(
            decisions_required_from_user=[
                DecisionRequired(
                    question="Approve the pilot batch?",
                    decision_class="red",
                    recommendation="Approve with the stated controls.",
                )
            ]
        ),
    )
    body = text_of(out.path)
    assert "Approve the pilot batch?" in body
    assert "RED" in body


def test_evidence_register_shows_reliability(tmp_path):
    out = build(
        tmp_path,
        evidence=[
            {
                "evidence_ref": "E-1",
                "reliability_level": 1,
                "source_title": "USP <711> Dissolution",
                "excerpt_or_locator": "Section 3",
                "originating_agent": "A07",
            }
        ],
    )
    body = text_of(out.path)
    assert "USP <711> Dissolution" in body
    assert "Authoritative" in body


def test_a_report_is_never_overwritten(tmp_path):
    build(tmp_path, version=1)
    with pytest.raises(FileExistsError):
        build(tmp_path, version=1)
    second = build(tmp_path, version=2)
    assert second.filename.endswith("_v2.docx")


def test_result_records_a_hash_for_the_audit_trail(tmp_path):
    out = build(tmp_path)
    assert len(out.sha256) == 64
    assert out.generator_version
