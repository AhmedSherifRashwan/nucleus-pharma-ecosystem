"""Report generation against real records (PRD s18, s32, Phase 12)."""

from __future__ import annotations

import uuid

import pytest
from docx import Document

from app.reports.service import ReportService
from app.schemas.executive import ExecutiveDecisionBrief
from app.storage.db import session_scope
from app.storage.repositories import (
    ArtifactRepository,
    AuditRepository,
    EvidenceRepository,
    GovernanceRepository,
    RunRepository,
)
from tests.integration.conftest import requires_db

pytestmark = [pytest.mark.asyncio, requires_db]


def brief(**kw) -> ExecutiveDecisionBrief:
    base = {
        "objective": "Assess replacing gelatin with HPMC.",
        "executive_conclusion": "Feasible with mitigations.",
        "recommendation": "Proceed to a pilot batch.",
        "confidence": "medium",
        "consulted_agents": ["A07"],
    }
    return ExecutiveDecisionBrief(**{**base, **kw})


async def _run() -> tuple[uuid.UUID, str]:
    async with session_scope() as s:
        run = await RunRepository(s).create(user_request="Assess the shell change.")
        return run.id, run.run_ref


def body_of(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


async def test_report_is_registered_as_a_versioned_artifact(tmp_path):
    run_id, run_ref = await _run()
    service = ReportService(output_dir=tmp_path)

    first, _ = await service.generate(run_id=run_id, run_ref=run_ref, brief=brief(), results=[])
    second, _ = await service.generate(run_id=run_id, run_ref=run_ref, brief=brief(), results=[])

    assert (first.version, second.version) == (1, 2)
    assert first.path.is_file() and second.path.is_file()
    # s18.3: the first version survives the second.
    assert first.path != second.path

    async with session_scope() as s:
        records = [
            a for a in await ArtifactRepository(s).list_for_run(run_id) if a.kind == "report"
        ]
    assert {r.version for r in records} == {1, 2}
    assert all(r.sha256 and r.generator_version for r in records)


async def test_evidence_and_risks_are_pulled_from_the_database(tmp_path):
    run_id, run_ref = await _run()
    async with session_scope() as s:
        await EvidenceRepository(s).create(
            source_type="regulatory",
            source_title="ICH Q1A(R2) stability guideline",
            reliability_level=1,
            run_id=run_id,
            originating_agent="A10",
            excerpt_or_locator="Section 2.1.7",
        )
        await GovernanceRepository(s).add_risk(
            run_id=run_id,
            description="Cross-linking may slow dissolution on storage.",
            likelihood="medium",
            impact="high",
            mitigation="Add enzymatic dissolution testing.",
        )

    result, warnings = await ReportService(output_dir=tmp_path).generate(
        run_id=run_id, run_ref=run_ref, brief=brief(), results=[]
    )
    body = body_of(result.path)
    assert "ICH Q1A(R2) stability guideline" in body
    assert "Cross-linking may slow dissolution on storage." in body
    assert "Add enzymatic dissolution testing." in body
    assert warnings == []


async def test_a_cited_evidence_id_that_does_not_exist_is_flagged(tmp_path):
    run_id, run_ref = await _run()
    _, warnings = await ReportService(output_dir=tmp_path).generate(
        run_id=run_id,
        run_ref=run_ref,
        brief=brief(key_evidence=["E-999 supports the shell change"]),
        results=[],
    )
    assert any("E-999" in w for w in warnings)


async def test_a_run_with_no_evidence_is_called_out(tmp_path):
    run_id, run_ref = await _run()
    _, warnings = await ReportService(output_dir=tmp_path).generate(
        run_id=run_id, run_ref=run_ref, brief=brief(), results=[]
    )
    assert any("No evidence items" in w for w in warnings)


async def test_generation_is_audited(tmp_path):
    run_id, run_ref = await _run()
    result, _ = await ReportService(output_dir=tmp_path).generate(
        run_id=run_id, run_ref=run_ref, brief=brief(), results=[]
    )
    async with session_scope() as s:
        events = await AuditRepository(s).for_run(run_id)
    generated = [e for e in events if e.event_type == "report.generated"]
    assert len(generated) == 1
    assert generated[0].event_metadata["sha256"] == result.sha256
