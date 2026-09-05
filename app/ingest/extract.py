"""Text extraction for project documents (PRD s14, s17.3).

Extraction is deliberately plain: the text a specialist cites has to be the text
that is in the document, so nothing here summarises, reformats or interprets. A
page or section number travels with every chunk, because an excerpt a reader
cannot locate in the original is not provenance.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from app.audit.logger import get_logger

log = get_logger(__name__)

SUPPORTED = {".pdf", ".docx", ".md", ".markdown", ".txt", ".csv", ".json", ".html", ".htm"}

# Big enough to carry an argument, small enough that a citation points somewhere
# specific. Overlap keeps a claim that straddles a boundary retrievable.
CHUNK_CHARS = 1800
CHUNK_OVERLAP = 200
# The smallest piece worth splitting at. Not a floor on what is kept: a short
# block often carries the identifiers that matter most — "Reference NDA: 207997",
# an assay result, a revision date — and dropping it loses exactly the line a
# specialist needed to cite.
MIN_SPLIT_CHARS = 80
MIN_KEEP_CHARS = 15


@dataclass
class ExtractedChunk:
    index: int
    content: str
    locator: str


@dataclass
class ExtractedDocument:
    title: str
    content_type: str
    text: str
    chunks: list[ExtractedChunk]
    page_count: int | None = None


class UnsupportedDocument(RuntimeError):
    """The file type cannot be read. Better said than silently skipped."""


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n")
    # Collapse the ragged whitespace PDF extraction leaves behind, without
    # joining paragraphs that were genuinely separate.
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _chunk(pages: list[tuple[str, str]]) -> list[ExtractedChunk]:
    """Chunk block by block, so each chunk keeps a locator a reader can follow."""
    chunks: list[ExtractedChunk] = []
    for locator, text in pages:
        text = _clean(text)
        if len(text) < MIN_KEEP_CHARS:
            continue
        if len(text) <= CHUNK_CHARS:
            chunks.append(ExtractedChunk(index=len(chunks), content=text, locator=locator))
            continue

        start = 0
        part = 0
        while start < len(text):
            end = min(start + CHUNK_CHARS, len(text))
            if end < len(text):
                # Prefer a paragraph or sentence boundary over a mid-word cut.
                window = text.rfind("\n\n", start + MIN_SPLIT_CHARS, end)
                if window == -1:
                    window = text.rfind(". ", start + MIN_SPLIT_CHARS, end)
                if window != -1:
                    end = window + 1
            body = text[start:end].strip()
            if len(body) >= MIN_KEEP_CHARS:
                chunks.append(
                    ExtractedChunk(
                        index=len(chunks), content=body, locator=f"{locator} (part {part + 1})"
                    )
                )
                part += 1
            start = end - CHUNK_OVERLAP if end < len(text) else end
    return chunks


def extract_pdf(path: Path) -> ExtractedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedDocument("pypdf is required to read PDF files") from exc

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # a damaged page must not lose the document
            log.warning("pdf_page_unreadable", path=path.name, page=number, error=str(exc)[:120])
            continue
        pages.append((f"page {number}", text))

    text = "\n\n".join(t for _, t in pages)
    if not _clean(text):
        raise UnsupportedDocument(
            f"{path.name} has no extractable text. It is probably a scan; run OCR first."
        )
    title = (reader.metadata.title or "").strip() if reader.metadata else ""
    return ExtractedDocument(
        title=title or path.stem,
        content_type="application/pdf",
        text=_clean(text),
        chunks=_chunk(pages),
        page_count=len(reader.pages),
    )


def extract_docx(path: Path) -> ExtractedDocument:
    from docx import Document as read_document

    document = read_document(str(path))
    blocks: list[tuple[str, str]] = []
    heading = "start of document"
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append((heading, "\n\n".join(buffer)))
            buffer.clear()

    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading") and paragraph.text.strip():
            flush()
            heading = paragraph.text.strip()[:120]
        elif paragraph.text.strip():
            buffer.append(paragraph.text.strip())
    flush()

    for number, table in enumerate(document.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            blocks.append((f"table {number}", "\n".join(rows)))

    text = "\n\n".join(t for _, t in blocks)
    if not _clean(text):
        raise UnsupportedDocument(f"{path.name} contains no readable text")
    return ExtractedDocument(
        title=path.stem,
        content_type=("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        text=_clean(text),
        chunks=_chunk(blocks),
    )


def extract_text(path: Path) -> ExtractedDocument:
    raw = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() in (".html", ".htm"):
        raw = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)

    blocks: list[tuple[str, str]] = []
    heading = "start of document"
    buffer: list[str] = []
    for line in raw.split("\n"):
        if line.startswith("#") and line.strip("# ").strip():
            if buffer:
                blocks.append((heading, "\n".join(buffer)))
                buffer = []
            heading = line.strip("# ").strip()[:120]
        else:
            buffer.append(line)
    if buffer:
        blocks.append((heading, "\n".join(buffer)))

    text = _clean(raw)
    if not text:
        raise UnsupportedDocument(f"{path.name} is empty")
    return ExtractedDocument(
        title=path.stem,
        content_type="text/plain",
        text=text,
        chunks=_chunk(blocks),
    )


def extract(path: Path) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedDocument(
            f"{path.name}: {suffix or 'no extension'} is not supported. "
            f"Supported: {', '.join(sorted(SUPPORTED))}"
        )
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    return extract_text(path)
